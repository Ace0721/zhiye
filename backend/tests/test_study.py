import json
from io import BytesIO

import httpx
from docx import Document

from app.config import Settings
from app.main import create_app
from app.models import AnswerRecord, Card, Deck
from app.services.study_material import build_study_docx
from app.schemas import StudyMaterial, StudyPoint


async def test_study_filters_deduplicates_exports_and_isolates_decks(config: Settings) -> None:
    application = create_app(config)
    async with application.router.lifespan_context(application):
        with application.state.sessions.begin() as session:
            deck = Deck(title="诉讼程序", source_filename="课堂.pptx")
            other = Deck(title="其他课件", source_filename="其他.pdf")
            empty = Deck(title="仅课程安排", source_filename="安排.pdf")
            session.add_all([deck, other, empty]); session.flush()
            cards = []
            for index, (deck_id, title, key) in enumerate([
                (deck.id, "课程学时安排", "共32学时"),
                (deck.id, "举证责任", "当事人对自己提出的主张提供证据"),
                (deck.id, "重复重点", "当事人对自己提出的主张提供证据。"),
                (other.id, "其他知识", "不应出现在第一份资料中"),
                (empty.id, "课程学时安排", "共32学时"),
            ]):
                card = Card(deck_id=deck_id, order_index=index, title=title, key_point=key,
                            plain_explanation="主张事实的一方需要拿出证据。", question="四选一题不应导出",
                            options_json=json.dumps(["a", "b", "c", "d"]), correct_index=0,
                            reexplain="例如主张借款存在，就应提供借条。")
                session.add(card); session.flush(); cards.append(card)
            session.add(AnswerRecord(card_id=cards[1].id, selected_index=1, is_correct=False))
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=application), base_url="http://test") as client:
            result = await client.get(f"/api/decks/{deck.id}/study")
            assert result.status_code == 200
            material = result.json()
            assert material["excluded_count"] == 2
            assert [point["title"] for point in material["points"]] == ["举证责任"]
            assert material["points"][0]["wrong_count"] == 1
            word = await client.get(f"/api/decks/{deck.id}/study.docx")
            assert word.status_code == 200
            assert "wordprocessingml.document" in word.headers["content-type"]
            assert "filename*=UTF-8''" in word.headers["content-disposition"]
            document = Document(BytesIO(word.content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            assert "重点速记" in text and "知识点简明解释" in text and "我的易错补充" in text
            assert "借条" in text and "答错 1 次" in text
            assert "课程学时" not in text and "其他知识" not in text and "四选一题" not in text
            assert document.paragraphs[0].style.name == "Title"
            for route in ("study", "study.docx"):
                assert (await client.get(f"/api/decks/999/{route}")).status_code == 404
                assert (await client.get(f"/api/decks/{empty.id}/{route}")).status_code == 422


def test_word_export_handles_xml_controls_and_no_wrong_answers() -> None:
    material = StudyMaterial(deck_id=1, title="复习\x00", excluded_count=0, points=[
        StudyPoint(card_id=1, title="知识", key_point="重点\x01", explanation="解释\x0b正常",
                   clarification="例子", wrong_count=0)])
    document = Document(BytesIO(build_study_docx(material)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "解释正常" in text
    assert "我的易错补充" not in text
