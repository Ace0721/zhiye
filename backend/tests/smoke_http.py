"""真实 HTTP 冒烟：独立后端进程 + 本地模拟 API + curl，不访问云端。"""

import argparse
import json
import os
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

from reportlab.pdfgen import canvas
from docx import Document


class TemporaryWorkspace(tempfile.TemporaryDirectory):
    def cleanup(self) -> None:
        # Windows 在终止进程树后可能稍晚才释放 SQLite 文件句柄。
        for attempt in range(30):
            try:
                super().cleanup()
                return
            except PermissionError:
                if attempt == 29:
                    raise
                time.sleep(0.1)


def free_port() -> int:
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        return listener.getsockname()[1]


class MockHandler(BaseHTTPRequestHandler):
    calls: list[str] = []
    def log_message(self, format: str, *args: Any) -> None:
        pass

    def do_POST(self) -> None:
        request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        assert self.path == "/v1/chat/completions"
        assert request["model"] == "smoke-model"
        time.sleep(0.1)
        marker = re.search(r"\[第 (\d+) 页/张\]", request["messages"][-1]["content"])
        refs = [{"page":int(marker.group(1)),"quote":"Newton second law F = ma; force equals mass times acceleration."}] if marker else []
        card = {
            "title": "牛顿第二定律", "plain_explanation": "力越大，加速度越大。质量越大，越难加速。",
            "key_point": "合力等于质量乘加速度", "question": "质量一定时合力加倍，加速度如何变化？",
            "options": ["加倍", "减半", "不变", "变为四倍"], "correct_index": 0,
            "reexplain": "推同一辆车，用两倍合力就获得两倍加速度。",
            "source_refs": refs,
        }
        prompt = request["messages"][0]["content"]
        self.calls.append("map" if "思维导图" in prompt else "notes" if "课程复习资料编辑" in prompt else "reduce" if "学科知识骨架" in prompt else "cards")
        content = json.dumps([card], ensure_ascii=False)
        if "思维导图" in prompt:
            content = json.dumps({"title": "牛顿运动规律", "overview": "合力和质量共同决定物体的加速度。", "branches": [
                {"label": "动力学关系", "relation": "遵循", "detail": "合外力决定加速度。", "source_refs":refs, "children": [
                    {"label": "F = ma", "relation": "表达为", "detail": "合力等于质量乘加速度，惯性参考系下成立。", "children": [], "source_refs":refs}]}]}, ensure_ascii=False)
        elif "课程复习资料编辑" in prompt:
            content = "# 牛顿第二定律复习\n理解合力、质量与加速度之间的关系。\n## 核心关系\n- F = ma；质量一定时，合力加倍，加速度加倍。\n"
            if refs:
                content += f"> 来源：第{refs[0]['page']}页/张｜{refs[0]['quote']}\n"
            content += "## 易混与边界\n- 公式中的力是合外力。"
        elif "学科知识骨架" in prompt:
            content = "牛顿第二定律：合外力F等于质量m乘加速度a；惯性参考系下适用。"
        data = json.dumps({
            "id": "smoke", "object": "chat.completion", "created": 0, "model": "smoke-model",
            "choices": [{"index": 0, "message": {"role": "assistant", "content": content},
                         "finish_reason": "stop"}],
        }, ensure_ascii=False).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


def main() -> None:
    """临时目录和进程均由脚本回收，真实密钥不参与测试。"""
    parser = argparse.ArgumentParser()
    parser.add_argument("--binary", type=Path)
    args = parser.parse_args()
    curl = shutil.which("curl.exe") or shutil.which("curl")
    if curl is None:
        raise RuntimeError("此验收脚本需要 curl")
    backend_port = free_port()
    mock = ThreadingHTTPServer(("127.0.0.1", 0), MockHandler)
    thread = threading.Thread(target=mock.serve_forever, daemon=True)
    thread.start()
    try:
        with TemporaryWorkspace(prefix="zhiye-smoke-") as temporary:
            root = Path(temporary)
            path = root / "20-pages.pdf"
            document = canvas.Canvas(str(path))
            for page in range(1, 21):
                for line in range(15):
                    document.drawString(40, 780 - line * 25, f"Page {page}, topic {line}: Newton second law F = ma; force equals mass times acceleration.")
                document.showPage()
            document.save()
            env = os.environ.copy()
            env.update(ZHIYE_PORT=str(backend_port), ZHIYE_DATA_DIR=str(root / "data"),
                       ZHIYE_AI_BASE_URL=f"http://127.0.0.1:{mock.server_port}/v1",
                       ZHIYE_AI_API_KEY="smoke-test-placeholder", ZHIYE_AI_MODEL="smoke-model",
                       ZHIYE_AI_TIMEOUT="10", ZHIYE_AI_CONCURRENCY="3", ZHIYE_TASK_TIMEOUT="30")
            command = [str(args.binary.resolve())] if args.binary else [sys.executable, "-m", "app.main"]
            error_path = root / "backend.log"
            error_stream = error_path.open("wb")
            process = subprocess.Popen(command,
                                       cwd=Path(__file__).resolve().parents[1], env=env,
                                       stdout=subprocess.DEVNULL, stderr=error_stream,
                                       start_new_session=os.name != "nt")

            def request(endpoint: str, *arguments: str) -> Any:
                result = subprocess.run([curl, "--fail-with-body", "--silent", "--show-error", "--max-time", "5",
                                         *arguments, f"http://127.0.0.1:{backend_port}{endpoint}"],
                                        capture_output=True, check=True)
                return json.loads(result.stdout)

            try:
                for _ in range(100):
                    if process.poll() is not None:
                        raise RuntimeError("后端进程启动失败")
                    try:
                        request("/api/health")
                        break
                    except subprocess.CalledProcessError:
                        time.sleep(0.1)
                else:
                    raise RuntimeError("后端健康检查超时")
                started = time.monotonic()
                uploaded = request("/api/upload", "-F", f"file=@{path}")
                states = []
                for _ in range(150):
                    state = request(f"/api/tasks/{uploaded['task_id']}")
                    states.append((state["status"], state["progress"], state["done_blocks"]))
                    if state["status"] in {"completed", "failed"}:
                        break
                    time.sleep(0.1)
                assert state["status"] == "completed", state
                cards = request(f"/api/decks/{state['deck_id']}/cards")
                assert state["done_blocks"] == state["total_blocks"] >= 3
                assert len(cards) == state["card_count"] > 0
                assert all(card["correct_index"] == 0 for card in cards)
                assert cards[0]["source_refs"]
                source = request(f"/api/decks/{state['deck_id']}/sources/{cards[0]['source_refs'][0]['page']}")
                assert "Newton second law" in source["text"]
                answer = request(f"/api/cards/{cards[0]['id']}/answer", "-H", "Content-Type: application/json", "-d", '{"selected_index":1}')
                assert answer["is_correct"] is False
                wrong = request(f"/api/decks/{state['deck_id']}/wrong")
                assert len(wrong) == 1
                summary = request(f"/api/decks/{state['deck_id']}/summary")
                assert summary["wrong_count"] == 1
                assert summary["wrong_cards"][0]["card_id"] == cards[0]["id"]
                summary_word = subprocess.run([curl, "--fail-with-body", "--silent", "--max-time", "10",
                    f"http://127.0.0.1:{backend_port}/api/decks/{state['deck_id']}/summary.docx"], capture_output=True, check=True)
                rescue = Document(BytesIO(summary_word.stdout))
                rescue_text = "\n".join(paragraph.text for paragraph in rescue.paragraphs)
                assert "考前救命包" in rescue_text and "正确答案：A. 加倍" in rescue_text
                material = request(f"/api/decks/{state['deck_id']}/study")
                assert material["points"]
                for kind in ("notes", "map"):
                    request(f"/api/decks/{state['deck_id']}/study/{kind}", "-X", "POST")
                    for _ in range(150):
                        generated = request(f"/api/decks/{state['deck_id']}/study/{kind}")
                        if generated["status"] in {"completed", "failed"}:
                            break
                        time.sleep(0.1)
                    assert generated["status"] == "completed", generated
                    assert generated["source_type"] == "course_text"
                exported = subprocess.run([curl, "--fail-with-body", "--silent", "--max-time", "10",
                                           f"http://127.0.0.1:{backend_port}/api/decks/{state['deck_id']}/study.docx"],
                                          capture_output=True, check=True)
                word = Document(BytesIO(exported.stdout))
                assert any(paragraph.text == "核心关系" for paragraph in word.paragraphs)
                settings = request("/api/settings")
                assert settings["api_key_configured"] and "api_key" not in settings
                assert request("/api/settings/test", "-X", "POST")["success"]
                prior_card_calls = MockHandler.calls.count("cards")
                selected = request("/api/upload", "-F", f"file=@{path}", "-F", "outputs=notes,map")
                for _ in range(150):
                    selected_state = request(f"/api/tasks/{selected['task_id']}")
                    if selected_state["status"] in {"completed", "failed"}:
                        break
                    time.sleep(0.1)
                assert selected_state["status"] == "completed", selected_state
                assert selected_state["generated_outputs"] == ["notes", "map"]
                assert selected_state["card_count"] == 0 and MockHandler.calls.count("cards") == prior_card_calls
                editable = {key: value for key, value in cards[0].items() if key not in {"id","deck_id","order_index","created_at"}}
                editable["title"] = "手动修订的标题"
                changed = request(f"/api/cards/{cards[0]['id']}", "-X", "PUT", "-H", "Content-Type: application/json", "-d", json.dumps(editable,ensure_ascii=False))
                assert changed["title"] == editable["title"] and changed["revision"] == 1
                learning_path = f"/api/decks/{state['deck_id']}/session"
                saved = request(learning_path, "-H", "Content-Type: application/json", "-d", '{}')
                saved = request(learning_path + "/all/answer", "-H", "Content-Type: application/json", "-d",
                                json.dumps(dict(revision=saved["revision"], index=0, selected_index=1)))
                assert len(saved["queue"]) == len(cards) + 1
                saved = request(learning_path + "/all", "-X", "PUT", "-H", "Content-Type: application/json", "-d",
                                json.dumps(dict(revision=saved["revision"], index=1)))
                restored = request(learning_path, "-H", "Content-Type: application/json", "-d", '{}')
                assert restored["index"] == 1 and restored["answers"] == saved["answers"]
                shelf = request('/api/library?offset_minutes=480')
                assert shelf["resume"]["index"] == 1 and shelf["stats"]["deck_count"] == 2
                print(json.dumps({"result": "PASS", "pages": 20, "blocks": state["total_blocks"],
                                  "cards": len(cards), "states": states,
                                  "frozen_binary": bool(args.binary), "review_export_settings": "PASS", "study_word": "PASS", "selection_edit_sources": "PASS", "learning_resume": "PASS",
                                  "elapsed_seconds": round(time.monotonic() - started, 2),
                                  "api": "local mock; no cloud performance claim"}))
            finally:
                # 单文件引导程序会再拉起子进程，必须清理整棵测试进程树。
                if os.name == "nt":
                    subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
                elif process.poll() is None:
                    os.killpg(process.pid, signal.SIGTERM)
                try:
                    process.communicate(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.communicate(timeout=5)
                error_stream.close()
                if sys.exc_info()[0] is not None:
                    # 仅用于本脚本的合成课件和占位密钥验收；后台日志不能堵住无人消费的 PIPE。
                    sys.stderr.write(error_path.read_text(encoding="utf-8", errors="replace")[-4000:])
    finally:
        mock.shutdown()
        mock.server_close()
        thread.join(timeout=5)


if __name__ == "__main__":
    main()
