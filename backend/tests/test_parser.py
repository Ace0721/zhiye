from pathlib import Path
from threading import Event

import pytest
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from app.services.parser import ParseError, parse_document, split_pages


def make_pdf(path: Path, pages: int = 20) -> None:
    document = canvas.Canvas(str(path))
    for number in range(1, pages + 1):
        document.drawString(40, 780, f"Page {number}: Newton second law F = ma.")
        for line in range(12):
            document.drawString(40, 750 - line * 25, f"Lesson {number}: force equals mass times acceleration. Example {line}.")
        document.showPage()
    document.save()


def test_pdf_twenty_pages_preserves_order(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pdf"
    make_pdf(path)
    blocks = parse_document(path)
    text = "\n".join(block.text for block in blocks)
    assert len(blocks) >= 3
    assert text.index("Page 1:") < text.index("Page 20:")
    assert all(len(block.text) <= 6000 for block in blocks)


def test_pptx_text_and_table(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Newton law"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "F = ma"
    path = tmp_path / "lecture.pptx"
    presentation.save(path)
    text = parse_document(path)[0].text
    assert "Newton law" in text and "F = ma" in text


def test_split_long_page_and_empty_pages() -> None:
    blocks = split_pages([(1, ""), (2, "A" * 250), (3, "last paragraph")], max_chars=100)
    assert len(blocks) >= 3
    assert all(len(block.text) <= 100 for block in blocks)
    assert [block.index for block in blocks] == list(range(len(blocks)))
    assert sum(block.text.count("A") for block in blocks) == 250
    assert "第 1 " not in blocks[0].source
    assert "last paragraph" in blocks[-1].text


def test_blank_pdf_and_corrupted_file(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    document = canvas.Canvas(str(path))
    document.showPage()
    document.save()
    with pytest.raises(ParseError, match="未提取到文字"):
        parse_document(path)
    path.write_bytes(b"not a pdf")
    with pytest.raises(ParseError, match="解析失败"):
        parse_document(path)


def test_parser_cancellation(tmp_path: Path) -> None:
    path = tmp_path / "lecture.pdf"
    make_pdf(path, 1)
    cancelled = Event()
    cancelled.set()
    with pytest.raises(ParseError, match="停止"):
        parse_document(path, cancelled)
