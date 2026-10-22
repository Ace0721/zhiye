import json
from io import BytesIO

import httpx
from docx import Document

from app.config import Settings
from app.main import create_app
from app.models import AnswerRecord, Card, Deck
from sqlalchemy import func, select


async def test_answers_wrong_dedup_summary_and_isolation(config: Settings) -> None:
    application = create_app(config)
    async with application.router.lifespan_context(application):
        with application.state.sessions.begin() as session:
            first, second = Deck(title="微积分", source_filename="a.pdf"), Deck(title="其他", source_filename="b.pdf")
            session.add_all([first, second]); session.flush()
            cards = [Card(deck_id=deck.id, order_index=index, title=f"知识点{index}", plain_explanation="解释",
                          key_point=f"重点{index}", question="问题", options_json=json.dumps(["a", "b", "c", "d"]),
                          correct_index=0, reexplain="举个例子") for index, deck in enumerate([first, first, second])]
            session.add_all(cards); session.flush()
            ids = [card.id for card in cards]
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=application), base_url="http://test") as client:
            for card_id, index in [(ids[0], 1), (ids[0], 2), (ids[0], 0), (ids[1], 3), (ids[2], 2)]:
                response = await client.post(f"/api/cards/{card_id}/answer", json={"selected_index": index})
                assert response.status_code == 200
                assert response.json()["is_correct"] is (index == 0)
            wrong = (await client.get(f"/api/decks/{first.id}/wrong")).json()
            assert [card["id"] for card in wrong] == ids[:2]
            summary = (await client.get(f"/api/decks/{first.id}/summary")).json()
            assert summary["wrong_count"] == 2 and summary["total_cards"] == 2
            assert [point["card_id"] for point in summary["points"]] == ids[:2]
            assert [point["wrong_count"] for point in summary["points"]] == [2, 1]
            assert "markdown" not in summary
            word = await client.get(f"/api/decks/{first.id}/summary.docx")
            assert word.status_code == 200
            assert "wordprocessingml.document" in word.headers["content-type"]
            assert "filename*=UTF-8''" in word.headers["content-disposition"]
            document = Document(BytesIO(word.content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            assert document.paragraphs[0].style.name == "Title"
            assert "答错 2 次" in text and "正确答案：A. a" in text and "举个例子" in text
            assert "知识点2" not in text
            assert text.index("知识点0") < text.index("知识点1")
            decks = (await client.get("/api/decks")).json()
            assert sorted(deck["card_count"] for deck in decks) == [1, 2]
            for value in [-1, 4, True, "1"]:
                assert (await client.post(f"/api/cards/{ids[0]}/answer", json={"selected_index": value})).status_code == 422
            assert (await client.post('/api/cards/999/answer', json={"selected_index": 0})).status_code == 404
            assert (await client.get('/api/decks/999/wrong')).status_code == 404
            assert (await client.get('/api/decks/999/summary')).status_code == 404
            assert (await client.get('/api/decks/999/summary.docx')).status_code == 404
        with application.state.sessions() as session:
            assert session.scalar(select(func.count(AnswerRecord.id))) == 5


async def test_summary_word_without_wrong_answers(config: Settings) -> None:
    application = create_app(config)
    async with application.router.lifespan_context(application):
        with application.state.sessions.begin() as session:
            deck = Deck(title="复习<正文>", source_filename="a.pdf")
            session.add(deck); session.flush()
            session.add(Card(deck_id=deck.id, order_index=0, title="定义", plain_explanation="解释",
                key_point="公式：a < b & b > 0", question="问题", options_json=json.dumps(["a", "b", "c", "d"]),
                correct_index=0, reexplain="例子"))
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=application), base_url="http://test") as client:
            response = await client.get(f"/api/decks/{deck.id}/summary.docx")
            assert response.status_code == 200
            document = Document(BytesIO(response.content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            assert "重点速记" in text and "暂无历史错题" in text
            assert "公式：a < b & b > 0" in text
            assert "高频易错重点" not in text
