"""本地抽取文字，保留页码并按段落限制块大小；不做 OCR。"""

from dataclasses import dataclass
from pathlib import Path
from threading import Event
from typing import Any

import pdfplumber
from pptx import Presentation


class ParseError(Exception):
    """可以直接向用户展示的解析错误。"""


@dataclass(frozen=True)
class TextBlock:
    index: int
    source: str
    text: str
    pages: dict[int, str] | None = None


def split_pages(pages: list[tuple[int, str]], max_chars: int = 6000) -> list[TextBlock]:
    """按原顺序合并短页，超长段落硬切分，块大小包含页码标记。"""
    if max_chars < 64:
        raise ValueError("分块上限不能小于 64 字符")
    blocks: list[TextBlock] = []
    parts: list[str] = []
    page_numbers: list[int] = []
    page_parts: dict[int, list[str]] = {}
    size = 0

    def flush() -> None:
        nonlocal size
        if parts:
            source = "第 " + ", ".join(map(str, dict.fromkeys(page_numbers))) + " 页/张"
            blocks.append(TextBlock(len(blocks), source, "\n".join(parts), {page: "\n".join(texts) for page, texts in page_parts.items()}))
            parts.clear()
            page_numbers.clear()
            page_parts.clear()
            size = 0

    for page_number, text in pages:
        marker = f"[第 {page_number} 页/张]\n"
        limit = max_chars - len(marker) - 1
        for paragraph in text.splitlines():
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            for offset in range(0, len(paragraph), limit):
                piece = marker + paragraph[offset:offset + limit]
                if size + len(piece) + 1 > max_chars:
                    flush()
                parts.append(piece)
                page_numbers.append(page_number)
                page_parts.setdefault(page_number, []).append(paragraph[offset:offset + limit])
                size += len(piece) + 1
    flush()
    return blocks


def parse_document(path: Path, cancel: Event | None = None) -> list[TextBlock]:
    """从 PDF 或 PPTX 提取文字；取消信号使线程在页边界尽快退出。"""
    pages: list[tuple[int, str]] = []

    def check_cancelled() -> None:
        if cancel is not None and cancel.is_set():
            raise ParseError("任务已停止")

    try:
        if path.suffix.lower() == ".pdf":
            with pdfplumber.open(path) as document:
                for number, page in enumerate(document.pages, 1):
                    check_cancelled()
                    pages.append((number, page.extract_text() or ""))
        elif path.suffix.lower() == ".pptx":
            # 显式关闭文件句柄，确保 Windows 下能删除临时文件。
            with path.open("rb") as stream:
                presentation = Presentation(stream)
            for number, slide in enumerate(presentation.slides, 1):
                check_cancelled()
                text: list[str] = []

                def collect(shapes: Any) -> None:
                    for shape in shapes:
                        if shape.has_text_frame:
                            text.append(shape.text_frame.text)
                        if shape.has_table:
                            text.extend(cell.text for row in shape.table.rows for cell in row.cells)
                        if hasattr(shape, "shapes"):
                            collect(shape.shapes)

                collect(slide.shapes)
                pages.append((number, "\n".join(text)))
        else:
            raise ParseError("仅支持 PDF 和 PPTX 文件")
    except ParseError:
        raise
    except Exception:
        raise ParseError("课件解析失败：文件可能损坏、加密，或内容与扩展名不符") from None
    check_cancelled()
    blocks = split_pages(pages)
    if not blocks:
        raise ParseError("未提取到文字：扫描版 PDF 和纯图片课件暂不支持，请使用含可选中文字的课件")
    return blocks
