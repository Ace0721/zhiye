"""用已提炼的课件知识生成复习资料，导出不再调用云端 AI。"""

import re
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import AnswerRecord, Card, Deck
from app.schemas import StudyMaterial, StudyPoint
from app.services.card_gen import is_course_admin_card


def build_study_material(session: Session, deck: Deck) -> StudyMaterial:
    """保持课件顺序，过滤课程事务与完全重复重点，并附上真实错题次数。"""
    cards = session.scalars(select(Card).where(Card.deck_id == deck.id).order_by(Card.order_index)).all()
    counts = dict(session.execute(select(AnswerRecord.card_id, func.count(AnswerRecord.id))
                                 .join(Card).where(Card.deck_id == deck.id, AnswerRecord.is_correct.is_(False))
                                 .group_by(AnswerRecord.card_id)).all())
    seen: set[str] = set()
    points: list[StudyPoint] = []
    for card in cards:
        identity = re.sub(r"\W+", "", card.key_point).casefold()
        if is_course_admin_card(card) or identity in seen:
            continue
        seen.add(identity)
        points.append(StudyPoint(card_id=card.id, title=card.title, key_point=card.key_point,
                                 explanation=card.plain_explanation, clarification=card.reexplain,
                                 wrong_count=counts.get(card.id, 0)))
    return StudyMaterial(deck_id=deck.id, title=deck.title, points=points,
                         excluded_count=len(cards) - len(points))


def document_text(text: str) -> str:
    """移除 XML 不允许的控制字符，避免异常课件导致 Word 导出失败。"""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)


def build_study_docx(material: StudyMaterial) -> bytes:
    """输出可编辑的 Word：速记清单、简明解释、个人错题补充，不重复选择题。"""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    document.styles["Title"].font.size = Pt(22)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 2"].font.size = Pt(12)
    document.add_paragraph(document_text(material.title) + " 精简复习资料", "Title")
    document.add_paragraph(f"本资料整理了课件已提炼的 {len(material.points)} 个知识点，先看速记清单，再查简明解释。"
                           "课程事务与重复重点已过滤，未作答也可使用。AI 内容请结合原课件核对。")
    document.add_heading("重点速记", level=1)
    for point in material.points:
        paragraph = document.add_paragraph()
        paragraph.add_run(document_text(point.title) + "：").bold = True
        paragraph.add_run(document_text(point.key_point))
    document.add_heading("知识点简明解释", level=1)
    for index, point in enumerate(material.points, 1):
        document.add_heading(f"{index} {document_text(point.title)}", level=2)
        document.add_paragraph(document_text(point.explanation))
    wrong = sorted((point for point in material.points if point.wrong_count), key=lambda point: -point.wrong_count)
    if wrong:
        document.add_heading("我的易错补充", level=1)
        for point in wrong:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{document_text(point.title)}（答错 {point.wrong_count} 次）：").bold = True
            paragraph.add_run(document_text(point.clarification))
    section.footer.paragraphs[0].text = "知页 Zhiye  课件复习资料"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
